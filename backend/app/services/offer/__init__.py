import io

from docx import Document
from docx.shared import Pt, RGBColor

from app.services.extraction import _libreoffice_to_pdf


def generate_offer_pdf(
    *,
    company: str,
    candidate_name: str,
    job_title: str,
    body_text: str,
    salary: str | None = None,
    start_date: str | None = None,
) -> bytes | None:
    """Render a branded offer letter to PDF (DOCX -> PDF via headless LibreOffice).

    Returns PDF bytes, or None if conversion failed.
    """
    doc = Document()

    title = doc.add_paragraph()
    run = title.add_run(company)
    run.bold = True
    run.font.size = Pt(20)
    run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    heading = doc.add_heading("Letter of Offer", level=1)
    heading.runs[0].font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    doc.add_paragraph(f"Dear {candidate_name},")
    for para in (body_text or "").split("\n"):
        doc.add_paragraph(para)

    if salary or start_date:
        doc.add_paragraph()
        details = doc.add_paragraph()
        details.add_run("Offer details").bold = True
        table = doc.add_table(rows=0, cols=2)
        table.style = "Light Grid Accent 1"
        rows = [("Position", job_title)]
        if salary:
            rows.append(("Compensation", salary))
        if start_date:
            rows.append(("Start date", start_date))
        for label, value in rows:
            cells = table.add_row().cells
            cells[0].text = label
            cells[1].text = str(value)

    doc.add_paragraph()
    doc.add_paragraph(f"Warm regards,\nThe {company} team")

    buf = io.BytesIO()
    doc.save(buf)
    return _libreoffice_to_pdf(buf.getvalue(), ".docx")
