import io
import subprocess
import tempfile
from pathlib import Path

from app.services.ai.base import ResumeDocument

IMAGE_MIMES = {"image/png", "image/jpeg", "image/jpg", "image/webp", "image/gif"}
PDF_MIME = "application/pdf"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
DOC_MIME = "application/msword"

EXT_TO_MIME = {
    ".pdf": PDF_MIME,
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".webp": "image/webp",
    ".gif": "image/gif",
    ".doc": DOC_MIME,
    ".docx": DOCX_MIME,
}

ALLOWED_EXTENSIONS = set(EXT_TO_MIME.keys())


def guess_mime(filename: str | None, fallback: str | None = None) -> str | None:
    if filename:
        ext = Path(filename).suffix.lower()
        if ext in EXT_TO_MIME:
            return EXT_TO_MIME[ext]
    return fallback


def _pdf_text(data: bytes) -> str:
    try:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        return "\n".join((page.extract_text() or "") for page in reader.pages).strip()
    except Exception:
        return ""


def _docx_text(data: bytes) -> str:
    try:
        from docx import Document

        doc = Document(io.BytesIO(data))
        return "\n".join(p.text for p in doc.paragraphs).strip()
    except Exception:
        return ""


def _libreoffice_to_pdf(data: bytes, suffix: str) -> bytes | None:
    """Convert a .doc/.docx to PDF using headless LibreOffice (installed in the worker image)."""
    try:
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / f"input{suffix}"
            src.write_bytes(data)
            subprocess.run(
                ["soffice", "--headless", "--convert-to", "pdf", "--outdir", tmp, str(src)],
                check=True,
                capture_output=True,
                timeout=120,
            )
            pdf = src.with_suffix(".pdf")
            if pdf.exists():
                return pdf.read_bytes()
    except Exception:
        return None
    return None


def normalize_document(data: bytes, filename: str | None, mime: str | None) -> ResumeDocument:
    """Turn raw uploaded bytes into a ResumeDocument the AI providers can consume.

    Strategy:
    - PDF / image  -> pass through natively (vision models read them directly).
    - DOCX         -> extract text; also try LibreOffice->PDF for layout fidelity.
    - DOC (legacy) -> LibreOffice->PDF (native) with no reliable text fallback.
    """
    resolved = guess_mime(filename, mime)

    if resolved == PDF_MIME:
        return ResumeDocument(
            data=data, media_type=PDF_MIME, filename=filename, text=_pdf_text(data) or None
        )

    if resolved in IMAGE_MIMES:
        return ResumeDocument(data=data, media_type=resolved, filename=filename)

    if resolved == DOCX_MIME:
        text = _docx_text(data)
        pdf = _libreoffice_to_pdf(data, ".docx")
        if pdf:
            return ResumeDocument(
                data=pdf, media_type=PDF_MIME, filename=filename, text=text or None
            )
        return ResumeDocument(text=text or "(empty document)", filename=filename)

    if resolved == DOC_MIME:
        pdf = _libreoffice_to_pdf(data, ".doc")
        if pdf:
            return ResumeDocument(data=pdf, media_type=PDF_MIME, filename=filename)
        return ResumeDocument(text="(could not read legacy .doc file)", filename=filename)

    # Unknown — best effort: treat as plain text.
    try:
        return ResumeDocument(text=data.decode("utf-8", errors="ignore"), filename=filename)
    except Exception:
        return ResumeDocument(text="(unsupported file type)", filename=filename)
